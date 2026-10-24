"""Build a one-page MVP Canvas using lecture 4, slides 13 and 39."""

from pathlib import Path
from html import escape
import os

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle


CANVAS_DATE = "24.10.2026"
LECTURE_URL = "https://drive.google.com/file/d/1GFWqP0PaKKxUEqsraXWDtrBXHeGtyCKg/view"
CARDS = [
    ("1. Пользователь", [
        "Игрок osu!standard, которому трудно выбрать карту и понять свои сильные и слабые навыки. Первая группа - игроки с местом в рейтинге от 10 000 до 9 999 999 (5-7 digit).",
        "Заказчик - Ляхов Я.П., автор идеи и руководитель проекта.",
    ]),
    ("2. Проблема и ценность", [
        "Общие звезды карты не учитывают навыки игрока. Карты с одинаковыми звездами могут требовать разного уровня aim и speed.",
        "Проверяем, помогает ли отчет заранее понять, что будет сложно, и осмысленнее выбрать карту для игры.",
    ]),
    ("3. Один сценарий", [
        "Ссылка на профиль -> загрузка доступной истории -> файл карты .osu/.osz -> персональный отчет.",
        "Игрок видит, какой навык создает основную сложность, шанс прохождения и ожидаемую точность. Новые игры после запуска для первого анализа не нужны.",
    ]),
    ("4. Данные", [
        "Публичный профиль: до 100 recent и 100 best через Hinamizawa; ответы в JSON. Повторы убираются, поэтому уникальных карт может быть меньше 200. Карта - .osu/.osz. Хранение - локально в SQLite.",
        "Для обучения нужны отдельные попытки с проходами и неудачами. Best не используем как полную историю для обучения шанса прохождения.",
    ]),
    ("5. Результат MVP", [
        "Локальное приложение с одним отчетом: профиль навыков aim/speed; сложность карты для игрока по этим навыкам; вероятность прохождения; ожидаемая accuracy при прохождении.",
        "В отчете видны объем истории и статус оценки. Aim - наведение на объекты; speed - скорость нажатий. При нехватке данных числовой прогноз не выдумываем.",
    ]),
    ("6. Три критерия готовности", [
        "AC-01: второй участник запускает приложение по README, загружает профиль и .osu/.osz без пароля игрока. Видны число результатов и выбранная карта. REQ-01, 02, 15.",
        "AC-02: для двух контрольных историй одна карта дает разные персональные оценки. Есть aim/speed, шанс 0..1, accuracy 0..100 и статус; при малой истории - сообщение о нехватке данных. REQ-03, 04, 05, 06, 09.",
        "AC-03: на новых попытках есть отчет: Brier score (ошибка вероятности), MAE (средняя ошибка accuracy на пройденных картах), сравнение с простым расчетом и число попыток. Модель хуже простого расчета не используется как проверенная. REQ-06, 10, 11.",
    ]),
    ("7. Ограничения", [
        "Только osu!standard, два навыка, локальный запуск и небольшой пилот. Публичная история неполная; источник данных и файлы карт могут быть недоступны.",
        "Подписка, платежи, другие режимы и рекомендации тренировочных карт по слабым навыкам отложены. Улучшение удержания игроков пока не проверяем.",
    ]),
    ("8. Следующее решение после демо", [
        "Заказчик решает, можно ли переходить к тестированию, что надо исправить и достаточно ли двух навыков. Основание - AC-01..03, отзывы игроков и список ошибок.",
        "Canvas описывает целевую версию к 14 ноября. Профиль и импорт уже есть в прототипе. Персональную сложность по навыкам и качество прогноза еще нужно проверить.",
    ]),
]


def build_canvas(team, deadline, output_dir):
    output_dir = Path(output_dir)
    cards = [(heading, [text.replace("Ляхов Я.П.", team["PM"]) for text in paragraphs])
             for heading, paragraphs in CARDS]
    subtitle = f"osu!standard | Версия {CANVAS_DATE} | Сдача {deadline:%d.%m.%Y}"
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(29.7), Cm(21)
    section.top_margin = section.bottom_margin = Cm(1.0)
    section.left_margin = section.right_margin = Cm(1.1)
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(9.5)
    style.paragraph_format.space_after = Pt(3)
    style.paragraph_format.line_spacing = 1.0
    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(2)
    run = title.add_run("pulse! - MVP Canvas")
    run.bold, run.font.size = True, Pt(17)
    doc.add_paragraph(subtitle)
    table = doc.add_table(rows=4, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    for row_index, row in enumerate(table.rows):
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for col_index, cell in enumerate(row.cells):
            cell.width = Cm(13.75)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            heading, paragraphs = cards[row_index * 2 + col_index]
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run(heading)
            r.bold, r.font.size = True, Pt(10)
            r.font.color.rgb = RGBColor.from_string("244A70")
            for text in paragraphs:
                cell.add_paragraph(text)
    footer = section.footer.paragraphs[0]
    footer.paragraph_format.space_after = Pt(0)
    r = footer.add_run("Основа: лекция 4, слайды 13 и 39. AC - критерий готовности; REQ - требование из реестра.")
    r.font.size = Pt(7)
    doc.save(output_dir / "MVP Canvas - pulse!.docx")

    font_dir = Path(os.environ.get("WINDIR", r"C:\Windows")) / "Fonts"
    pdfmetrics.registerFont(TTFont("CanvasArial", str(font_dir / "arial.ttf")))
    pdfmetrics.registerFont(TTFont("CanvasArialBold", str(font_dir / "arialbd.ttf")))
    body = ParagraphStyle("CanvasBody", fontName="CanvasArial", fontSize=9.5, leading=11.4, spaceAfter=4)
    heading_style = ParagraphStyle("CanvasHeading", fontName="CanvasArialBold", fontSize=10, leading=12,
                                   textColor=colors.HexColor("#244A70"), spaceAfter=5)
    title_style = ParagraphStyle("CanvasTitle", fontName="CanvasArialBold", fontSize=17, leading=20, spaceAfter=3)
    subtitle_style = ParagraphStyle("CanvasSubtitle", fontName="CanvasArial", fontSize=8.5, leading=11, spaceAfter=7)
    grid = []
    for i in range(0, 8, 2):
        cells = []
        for heading, paragraphs in cards[i:i + 2]:
            cells.append([Paragraph(escape(heading), heading_style)] +
                         [Paragraph(escape(text), body) for text in paragraphs])
        grid.append(cells)
    page_width, page_height = landscape(A4)
    margin = 29
    pdf_table = Table(grid, colWidths=[(page_width - 2 * margin) / 2] * 2)
    pdf_table.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#AABAC8")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 8),
        ("RIGHTPADDING", (0, 0), (-1, -1), 8),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    def footer_pdf(canvas, document):
        canvas.setFont("CanvasArial", 7)
        canvas.setFillColor(colors.HexColor("#555555"))
        canvas.drawString(margin, 17, "Основа: лекция 4, слайды 13 и 39. AC - критерий готовности; REQ - требование из реестра.")
        canvas.linkURL(LECTURE_URL, (margin, 14, page_width - margin, 25), relative=0)
    pdf = SimpleDocTemplate(str(output_dir / "MVP Canvas - pulse!.pdf"), pagesize=landscape(A4),
                            leftMargin=margin, rightMargin=margin, topMargin=20, bottomMargin=30)
    pdf.build([Paragraph("pulse! - MVP Canvas", title_style), Paragraph(escape(subtitle), subtitle_style), pdf_table],
              onFirstPage=footer_pdf, onLaterPages=footer_pdf)


if __name__ == "__main__":
    from build_course_pack import TEAM, DEADLINE, ROOT
    build_canvas(TEAM, DEADLINE, ROOT)
    print("Created MVP Canvas DOCX and PDF")
