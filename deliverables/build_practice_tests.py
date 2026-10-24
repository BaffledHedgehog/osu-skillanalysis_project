"""Build the five test cases requested in practice 1."""
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUTPUT = Path(__file__).parent / 'Практика 1 - Тест-кейсы MVP - pulse!.docx'

CASES = [
    ('TC-01', 'REQ-02', 'Загрузка корректной карты', 'Успешный',
     'MVP запущен. Подготовлены корректная карта osu!standard .osu и .osz с этой картой. Известны ее ID, название и число объектов.',
     '1. Загрузить файл .osu.\n2. Открыть анализ и сравнить данные с исходным файлом.\n3. Загрузить .osz с той же картой и повторить проверку.',
     'Оба формата приняты. ID, название и число объектов совпадают с исходной картой. Stars, aim и speed рассчитаны. Ошибки нет.',
     'Исходные .osu и .osz, скриншот анализа, ответы API загрузки и анализа.'),
    ('TC-02', 'REQ-06', 'Модель лучше базового прогноза', 'Успешный',
     'Тестовый стенд; достаточно размеченных попыток для обучения. Проверочные карты исключены из обучения. Контрольные значения: Brier модели 0.12, базового прогноза 0.20. Карта запроса входит в поддерживаемый диапазон.',
     '1. В автотесте задать контрольные результаты расчета Brier.\n2. Запустить проверку модели.\n3. Запросить прогноз карты и проверить выбранную модель.',
     'Модель принята: статус trained. В прогнозе используется ML прохождения. Метрики показывают 0.12 и 0.20; интерфейс обозначает прогноз как ML.',
     'Отчет автотеста с контрольными метриками, ответ API проверки модели и прогноза, скриншот статуса.'),
    ('TC-03', 'REQ-02', 'Поврежденный файл карты', 'Негативный',
     'MVP запущен. Список ID карт в каталоге сохранен. Подготовлен обычный текстовый файл без структуры карты, переименованный в .osu.',
     '1. Загрузить поврежденный файл.\n2. Проверить сообщение об ошибке.\n3. Обновить каталог и сравнить список ID с исходным.',
     'Показана понятная ошибка о неверном файле карты. Новая карта не добавлена. Приложение продолжает отвечать на запросы.',
     'Поврежденный файл, скриншот ошибки, ответ API загрузки, списки ID каталога до и после.'),
    ('TC-04', 'REQ-06', 'Модель хуже базового прогноза', 'Негативный',
     'Тестовый стенд; достаточно размеченных попыток. Нет пересечения карт обучения и проверки. Контрольные значения: Brier модели 0.28, базового прогноза 0.20. Карта запроса входит в поддерживаемый диапазон.',
     '1. В автотесте задать контрольные результаты расчета Brier.\n2. Запустить проверку модели.\n3. Запросить прогноз карты и проверить его источник.',
     'Модель отклонена: статус rejected. Для карты используется предварительный расчет. Интерфейс не обозначает его как проверенный ML-прогноз.',
     'Отчет автотеста, ответы API с метриками и источником прогноза, скриншот статуса.'),
    ('TC-05', 'REQ-06', 'Ошибка модели равна базовой', 'Граничный',
     'Тестовый стенд; достаточно размеченных попыток. Нет пересечения карт обучения и проверки. Контрольные значения Brier модели и базового прогноза равны: 0.20. Карта запроса входит в поддерживаемый диапазон.',
     '1. В автотесте задать одинаковые значения Brier: 0.20 и 0.20.\n2. Запустить проверку модели.\n3. Запросить прогноз карты и проверить решение.',
     'Модель отклонена: статус rejected. Равенство не считается улучшением. Используется предварительный расчет, статус trained не присваивается.',
     'Отчет граничного автотеста, ответы API с равными метриками и источником прогноза.'),
]


def table(doc, fields):
    result = doc.add_table(rows=0, cols=2)
    result.style = 'Table Grid'
    result.autofit = False
    result.columns[0].width = Cm(3.8)
    result.columns[1].width = Cm(13.2)
    for label, text in fields:
        cells = result.add_row().cells
        cells[0].text, cells[1].text = label, text
        cells[0].width, cells[1].width = Cm(3.8), Cm(13.2)
        for run in cells[0].paragraphs[0].runs:
            run.bold = True
        shade = OxmlElement('w:shd')
        shade.set(qn('w:fill'), 'F1EDF8')
        cells[0]._tc.get_or_add_tcPr().append(shade)
        for cell in cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(3)
                paragraph.paragraph_format.space_before = Pt(3)
                for run in paragraph.runs:
                    run.font.size = Pt(10)
        row_properties = cells[0]._tc.getparent().get_or_add_trPr()
        row_properties.append(OxmlElement('w:cantSplit'))
    return result


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.left_margin = section.right_margin = Cm(2)
    section.top_margin = section.bottom_margin = Cm(1.6)
    section.footer_distance = Cm(.8)
    normal = doc.styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.05
    language = OxmlElement('w:lang')
    language.set(qn('w:val'), 'ru-RU')
    normal.element.get_or_add_rPr().append(language)
    for style in ('Title', 'Heading 1', 'Heading 2'):
        doc.styles[style].font.name = 'Calibri'
        doc.styles[style].font.color.rgb = RGBColor.from_string('54259B')
    doc.styles['Title'].font.size = Pt(22)
    doc.styles['Heading 1'].font.size = Pt(15)
    doc.styles['Heading 2'].font.size = Pt(12)

    doc.core_properties.title = 'Практика 1. Тест-кейсы MVP pulse!'
    doc.core_properties.subject = 'Управление ИТ-проектами. Проектирование тестов для MVP.'
    doc.core_properties.author = 'Команда проекта pulse!'
    doc.add_heading('Практика 1. Тест-кейсы MVP', 0)
    doc.add_paragraph('pulse! - анализ навыков и персональной сложности карт osu!standard')
    doc.add_paragraph('Университет Иннополис | Управление ИТ-проектами | 24.10.2026')
    doc.add_heading('Требования и критерии приемки', 1)
    doc.add_paragraph('REQ-02, Must. Приложение принимает .osu и .osz карты osu!standard и извлекает данные карты. Приемка: оба формата загружаются, режим, ID и признаки определены верно. Поврежденный файл отклоняется с понятной ошибкой.')
    doc.add_paragraph('REQ-06, Must. Обученная модель прохождения используется только при Brier score ниже базового прогноза на проверочных картах. Приемка: более точная модель применяется; при худшей или равной ошибке показывается предварительный расчет. Brier score - ошибка вероятности; меньше означает лучше.')
    doc.add_paragraph('Объект проверки: рабочая версия MVP от 24.10.2026. Ответственный за тестирование - Темиров А.М. Исправления импорта - Гайфуллин А.Р.; модели и метрики - Недугов Д.')
    doc.add_paragraph('Всего 5 кейсов: 2 успешных, 2 негативных, 1 граничный. Значения Brier - контрольные данные автотестов, а не метрики качества на реальных игроках.')
    for index, (tc, req, title, kind, preconditions, steps, expected, evidence) in enumerate(CASES):
        if index in (1, 3):
            doc.add_page_break()
        doc.add_heading(f'{tc}. {title}', 2)
        table(doc, [
            ('TC-ID / REQ-ID', f'{tc} / {req}'), ('Тип', kind),
            ('Предусловия', preconditions), ('Шаги', steps),
            ('Ожидаемый результат', expected), ('Ответственный', 'Темиров А.М., тестирование'),
            ('Evidence', evidence), ('Фактический результат', 'Заполняется после выполнения.'),
            ('Статус', 'Не выполнен. После проверки: Pass / Fail / Blocked.'),
        ])
        doc.add_paragraph()
    doc.add_paragraph('При Fail сохраняем описание ошибки и evidence. При Blocked записываем причину, которая помешала выполнить проверку.')
    footer = section.footer.paragraphs[0]
    footer.add_run('pulse! | Практика 1 | ')
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    for run in footer.runs:
        run.font.size = Pt(9)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == '__main__':
    build()
